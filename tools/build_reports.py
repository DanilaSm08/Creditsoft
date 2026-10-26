"""Build four evidence-based Russian DOCX reports using the Lab 1–4 cover style."""

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
EVIDENCE = ROOT / "docs/evidence"
REPORTS = ROOT / "reports"


def read(name):
    return json.loads((EVIDENCE / name).read_text(encoding="utf-8"))


class Report:
    def __init__(self, number, title):
        self.number = number
        self.document = Document()
        self.figures = 0
        self.tables = 0
        doc = self.document
        section = doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
        section.different_first_page_header_footer = True
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(14)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.first_line_indent = Cm(1.25)
        for name, size in [("Heading 1", 16), ("Heading 2", 14)]:
            style = doc.styles[name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.first_line_indent = Cm(0)
            style.paragraph_format.keep_with_next = True
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Cm(0)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        cover = [
            "Министерство науки и высшего образования Российской Федерации",
            "Федеральное государственное бюджетное образовательное учреждение высшего образования",
            "«УЛЬЯНОВСКИЙ ГОСУДАРСТВЕННЫЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ»",
            "",
            "Факультет информационных систем и технологий",
            "Кафедра «Информационные системы»",
            "",
            "Дисциплина «Конструирование программного обеспечения»",
            "",
            "",
            f"Лабораторная работа №{number}",
            title,
            "",
            "",
        ]
        for text in cover:
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1
        for text in [
            "Выполнил:",
            "студент гр. ПИбд-42",
            "Смолин Д.С.",
            "",
            "Проверил:",
            "Романов А. А.",
        ]:
            p = doc.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
        p = doc.add_paragraph("Ульяновск, 2026 г.")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(28)
        doc.add_page_break()

    def h(self, text):
        self.document.add_heading(re.sub(r"^(\d+) ", r"\1. ", text), level=1)

    def p(self, text):
        p = self.document.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def table(self, caption, headers, rows, widths=None):
        self.tables += 1
        p = self.document.add_paragraph(f"Таблица {self.tables} — {caption}")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        table = self.document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = False
        if widths:
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Cm(width)
        header = table.rows[0]
        repeat = OxmlElement("w:tblHeader")
        header._tr.get_or_add_trPr().append(repeat)
        for cell, value in zip(header.cells, headers, strict=True):
            cell.text = str(value)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EFF3")
            cell._tc.get_or_add_tcPr().append(shade)
        for values in rows:
            for cell, value in zip(table.add_row().cells, values, strict=True):
                cell.text = str(value)
        for row_index, row in enumerate(table.rows):
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for index, cell in enumerate(row.cells):
                if widths:
                    cell.width = Cm(widths[index])
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(11)
                        run.bold = row_index == 0
        self.document.add_paragraph().paragraph_format.space_after = Pt(0)

    def figure(self, path, caption):
        path = ROOT / path
        if not path.exists():
            raise FileNotFoundError(path)
        self.figures += 1
        with Image.open(path) as img:
            width, height = img.size
        scale = min(16 / width, 15.5 / height)
        p = self.document.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width * scale))
        p = self.document.add_paragraph(f"Рисунок {self.figures} — {caption}")
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(11)

    def finish(self):
        self.h("Источники и воспроизведение")
        self.p(
            "Я использовал задания лабораторных работ №5–8, план задач и оценки из ЛР3, архитектуру из ЛР4. Отчёты Lab_1.docx–Lab_4.docx сохранил без изменений."
        )
        self.p(
            "Я разместил исходный код в D:\\STUDY\\КПО\\CreditSoft и в репозитории https://github.com/DanilaSm08/Creditsoft. Приложение доступно по http://127.0.0.1:8080/. Инструкции запуска находятся в README.md, результаты проверок — в docs/evidence и artifacts/runs, скриншоты — в docs/images."
        )
        self.p(
            "Документация инструментов: https://docs.djangoproject.com/en/5.2/; https://docs.pytest.org/; https://playwright.dev/python/; https://docs.github.com/actions; https://docs.astral.sh/ruff/; https://www.conventionalcommits.org/ru/v1.0.0/."
        )
        REPORTS.mkdir(exist_ok=True)
        path = REPORTS / f"Lab_{self.number}.docx"
        self.document.save(path)
        shutil.copy2(path, ROOT.parent / path.name)
        print(f"{path.name}: {self.figures} figures, {self.tables} tables")


def lab5(summary, tasks):
    r = Report(5, "Управление исходным кодом и отслеживание метрик")
    r.h("1 Цель и исходные условия")
    r.p(
        "Я создал репозиторий CreditSoft и настроил контроль изменений, тестирование, сборку и измерение качества. Система рассчитывает кредитоспособность физического лица по анкете и подбирает предложения из каталога."
    )
    r.p(
        "Я разместил проект на GitHub в репозитории DanilaSm08/Creditsoft. Использовал Issues для задач, Pull requests для объединения изменений, GitHub Actions для проверок и Releases для публикации версий и отчётов."
    )
    r.figure(
        "docs/images/infrastructure/repository.png",
        "Репозиторий CreditSoft на GitHub",
    )
    r.h("2 Структура репозитория и правила изменений")
    r.table(
        "Назначение основных каталогов",
        ["Путь", "Содержимое"],
        [
            ("core/", "Модели, формы, прикладная логика, шаблоны и статические ресурсы"),
            ("creditsoft/", "Настройки Django, маршрутизация и WSGI"),
            ("tests/", "Модульные, интеграционные и браузерные проверки"),
            (
                ".github/workflows/",
                "Проверки качества и сборка в GitHub Actions",
            ),
            ("tools/; .githooks/", "Сборка, метрики, отчёты, работа с GitHub; проверка сообщений"),
            (
                "infra/; Dockerfile; compose.yaml",
                "Nginx, приложение и PostgreSQL",
            ),
            (
                "docs/; reports/; artifacts/",
                "Документация, доказательства, Word/PDF, полные результаты проверок",
            ),
        ],
        [5.2, 11.8],
    )
    r.p(
        "Я организовал работу в коротких ветках: main содержит объединённую версию, feature/prototype — прототип, feature/assessment — расчётный контур, release/1.0.0 — выпуск. В ci/protection-probe я добавил намеренно падающий тест для проверки защиты. Изменения проходят локальные проверки, push, pull request и GitHub Actions. Секреты и окружение исключены через .gitignore."
    )
    r.p(
        "Я использовал Conventional Commits: type(scope): description или type: description. Пример: feat: реализовал анкету заёмщика. Hook commit-msg запускает tools/commit_message.py, а CI повторно проверяет сообщения всей истории."
    )
    r.table(
        "Проверенные ограничения",
        ["Проверка", "Результат и доказательство"],
        [
            (
                "Некорректное сообщение invalid message",
                "Отклонено, код возврата 1; protection-proof.json",
            ),
            (
                "Прямой push в main",
                "Отклонён; правило запрещает push и force-push; protection-proof.json",
            ),
            (
                "PR №32 с намеренно падающим тестом",
                "Слияние отклонено при ошибке обязательной проверки quality; PR закрыт без слияния",
            ),
            (
                "PR №31 / №33",
                "Я объединил прототип и основной функционал через pull requests",
            ),
            (
                "Обход администратором",
                "enforce_admins=true; правила защиты действуют для администратора",
            ),
        ],
        [6, 11],
    )
    r.p(
        "Я настроил обязательную проверку quality и обновление ветки относительно main перед слиянием. Защита действует и для администратора. Прямое внесение изменений, принудительная отправка и удаление main запрещены."
    )
    r.figure("docs/images/infrastructure/branches.png", "Правила защиты main: ответ GitHub API")
    r.figure(
        "docs/images/infrastructure/merged-prototype.png",
        "Слияние прототипа через pull request №31",
    )
    r.h("3 Pipeline качества и сборки")
    r.figure(
        "docs/images/infrastructure/blocked-pr.png",
        "Контрольный pull request №32 с ошибкой обязательной проверки",
    )
    r.p(
        "Я настроил workflow CreditSoft CI на события push и pull_request. Он получает исходники, устанавливает Python 3.12, зависимости и Chromium, проверяет коммиты и запускает tools/check.py. Скрипт последовательно выполняет Ruff, проверку форматирования и миграций, pytest с покрытием и сборку wheel/sdist. При ошибке проверок сборка останавливается. Результаты сохраняются в артефакт quality-evidence."
    )
    r.table(
        "Автоматические результаты",
        ["Проверка", "Артефакт / условие"],
        [
            ("Ruff check + format", "lint.json, format.log; 0 замечаний и код завершения 0"),
            ("Django migrations", "migrations.log; makemigrations --check --dry-run"),
            (
                "pytest + Chromium",
                "junit.xml, tests.html, tests.log, test-map.json; все тесты успешны",
            ),
            (
                "coverage.py",
                "coverage.json/xml, HTML; общий порог 85%, строки/ветви показаны отдельно",
            ),
            (
                "Python build",
                "dist/creditsoft-<version>.whl и .tar.gz только после успешных проверок",
            ),
            ("Radon + таймеры", "metrics.json и docs/evidence/metrics/<этап>.json"),
        ],
        [5.5, 11.5],
    )
    r.p(
        "Я настроил GitHub Actions с отдельным сервисом PostgreSQL 16. Развёртывание приложения выполняю через Docker Compose: Nginx принимает запросы, Gunicorn запускает Django, PostgreSQL хранит данные. Работу этого контура я проверил командой smoke."
    )
    r.figure(
        "docs/images/infrastructure/actions.png",
        "Запуски GitHub Actions",
    )
    r.h("4 Метрики этапов")
    stages = [
        "baseline",
        "prototype-browser",
        "assessment-browser",
        "postgres-regression",
        "release",
        "github-ci",
    ]
    metrics = [
        read(f"metrics/{name}.json")
        for name in stages
        if (EVIDENCE / f"metrics/{name}.json").exists()
    ]
    r.table(
        "Автоматически измеренные показатели качества",
        ["Этап", "Тесты", "Строки, %", "Ветви, %", "Ruff", "Тесты, с"],
        [
            (
                m["stage"],
                m["tests"]["passed"],
                m["line_coverage_percent"],
                "н/п" if m["stage"] == "baseline" else m["branch_coverage_percent"],
                m["lint_issues"],
                m["test_duration_seconds"],
            )
            for m in metrics
        ],
        [4.4, 1.7, 2.6, 2.4, 1.6, 4.3],
    )
    r.table(
        "Сложность, объём и длительность",
        ["Этап", "CC ср./макс.", "Python-файлы", "SLOC", "Pipeline, с"],
        [
            (
                m["stage"],
                f"{m['mean_complexity']} / {m['max_complexity']}",
                m["source_files"],
                m["source_lines"],
                m["pipeline_seconds"],
            )
            for m in metrics
        ],
        [4.4, 3.1, 3.2, 2, 4.3],
    )
    r.p(
        "Я собрал показатели Ruff, строкового и ветвевого покрытия, числа тестов, длительности тестов и pipeline, средней и максимальной цикломатической сложности, объёма исходников. SLOC и Radon рассчитаны для Python-модулей core и creditsoft без миграций. Покрытие рассчитано для core без миграций и management."
    )
    r.p(
        "Я сохранил замеры каждого этапа в docs/evidence/metrics. Поле pipeline_seconds содержит длительность tools/check.py. Каждый запуск GitHub Actions связан с SHA проверенного коммита."
    )
    r.h("5 Выпуск версии и вывод")
    r.p(
        "Я отметил прототип аннотированным тегом v0.1.0, завершённую версию — v1.0.0. В GitHub Releases разместил пакеты wheel/sdist, отчёты, архивы исходников и результатов проверок, файл SHA256SUMS.txt. Команды выпуска и восстановления описал в docs/RELEASE.md."
    )
    r.figure("docs/images/infrastructure/releases.png", "Версии проекта в GitHub")
    r.finish()


def lab6(summary, tasks):
    r = Report(6, "Прототипирование")
    r.h("1 Цель и выбранные задачи")
    r.p(
        "Я выбрал вертикальный срез: регистрация → вход → анкета → проверка полей → сохранение → повторное открытие. На нём проверил взаимодействие Django Auth, ORM, шаблонов и Bootstrap с базой данных."
    )
    selected = [
        t
        for t in tasks
        if t["id"] == "PR-01" or t["id"].startswith(tuple(f"CR-0{n}" for n in range(1, 5)))
    ]
    r.table(
        "Состав прототипа по исходному плану",
        ["ID", "Задача", "План, ч"],
        [(t["id"], t["title"], t["planned_hours"]) for t in selected],
        [2.4, 12.3, 2.3],
    )
    r.p(
        f"Всего выбрано {len(selected)} задач, сумма исходных оценок — {sum(t['planned_hours'] for t in selected):g} ч. Реализованный прототип сохранён в теге v0.1.0 и слит через PR №31. Его можно изучать независимо от последующего расширения приложения."
    )
    r.h("2 Реализация и результаты")
    r.p(
        "Я использовал Django Auth для регистрации, проверки пароля, хеширования и входа по сессии. Модель BorrowerProfile хранит сведения о заёмщике, FinancialData — финансовые значения. Владелец определяется по request.user. Я проверяю обе формы до записи и сохраняю их одной транзакцией."
    )
    r.figure("docs/images/prototype/login-desktop.png", "Экран входа работающего прототипа")
    r.figure(
        "docs/images/prototype/questionnaire-desktop.png",
        "Сохранённая анкета с персональной и финансовой частями",
    )
    r.p(
        "Я настроил проверку обязательных полей, формата паспорта, возраста 18–75 лет, денежных сумм, срока кредита и стажа. Ошибки показываю рядом с полями. Текущие платежи могут превышать доход: такие значения используются при оценке высокой долговой нагрузки."
    )
    r.figure(
        "docs/images/prototype/validation-desktop.png",
        "Ошибки серверной валидации без частичной записи данных",
    )
    r.figure(
        "docs/images/prototype/questionnaire-mobile.png",
        "Анкета прототипа при ширине экрана 390 px",
    )
    r.h("3 Проверка прототипа")
    m = read("metrics/prototype-browser.json")
    r.table(
        "Фактические результаты этапа",
        ["Показатель", "Результат"],
        [
            ("Тесты", f"{m['tests']['passed']} успешных, включая браузерный сценарий"),
            ("Покрытие измеряемого Python-кода", "100% строк / 100% ветвей"),
            ("Линтер", "0 замечаний Ruff"),
            ("Время тестов", f"{m['test_duration_seconds']} с"),
            (
                "Пользовательский сценарий",
                "Регистрация, вход, неверные данные, сохранение, перезагрузка, мобильная ширина",
            ),
            ("Версия", "v0.1.0; docs/evidence/metrics/prototype-browser.json"),
        ],
        [6, 11],
    )
    r.h("4 Уточнение плана и оценок")
    r.p(
        "По результатам прототипа я выбрал модульный монолит. Django обеспечивает сессии, CSRF, хеширование, формы и транзакции. Я уточнил атомарное сохранение двух форм, проверку адреса перенаправления после входа, границы стажа и отображение на мобильном экране."
    )
    r.p(
        "В следующем инкременте я реализовал CR-05–CR-08: долговую нагрузку, рейтинг, каталог, фильтрацию и итоговый экран. Добавил сохранение снимка входных параметров и предложений, чтобы изменения анкеты и каталога сохраняли предыдущие результаты."
    )
    phase = next(p for p in summary["worklog"] if p["phase"] == "prototype")
    r.p(
        f"Я зафиксировал интервал этапа прототипирования: {phase['elapsed_seconds'] / 60:.2f} мин. В план этого этапа вошли 13 задач общей оценкой 62 часа. Выводы по прототипу использовал при реализации расчётного контура."
    )
    r.finish()


def lab7(summary, tasks):
    r = Report(7, "Инфраструктура разработки проекта")
    r.h("1 Среда автоматического тестирования")
    r.p(
        "Я настроил pytest, pytest-django, pytest-cov, pytest-html и pytest-playwright. Закрепил зависимости в requirements-dev.txt, параметры — в pyproject.toml. Локальные проверки выполняю на SQLite, интеграционные — на PostgreSQL 16. Браузерные сценарии запускаю в Chromium через Django test server. GitHub Actions использует отдельную тестовую базу."
    )
    r.table(
        "Уровни проверок",
        ["Уровень", "Что подтверждается"],
        [
            (
                "Модульный",
                "Чистые Decimal-вычисления, граничные баллы, нагрузка, валидация, критерии предложений",
            ),
            (
                "Интеграционный Django + БД",
                "Auth/CSRF, формы и транзакции, фильтрация, снимки результатов, права владельца, повторный расчёт",
            ),
            (
                "Пользовательский Chromium",
                "Полный путь пользователя и отрицательный сценарий; мобильная вёрстка, логотипы, пустые предложения",
            ),
            (
                "Развёрнутый стенд",
                "Nginx → Gunicorn → PostgreSQL, health, вход, анкета, рейтинг 900 и четыре предложения, перезагрузка",
            ),
        ],
        [5, 12],
    )
    r.h("2 Проектирование тестов и ожидаемые результаты")
    r.table(
        "Значимые проверочные случаи",
        ["Вход / действие", "Ожидаемый результат"],
        [
            (
                "Доход 100 000; платежи 15 000; сумма 300 000; стаж 48 месяцев",
                "Нагрузка 15%; баллы 350 + 250 + 300 = 900; 4 предложения по возрастанию ставки",
            ),
            (
                "Нагрузка на границах 10, 30, 50, 70% и сразу выше",
                "Переходы 450 → 350 → 200 → 80 → 0; округление отображения не меняет ветвь расчёта",
            ),
            ("Стаж 0/1/6/12/36 месяцев", "Баллы 0/40/100/180/250; границы включительны"),
            ("Сумма / годовой доход на границах 0.5/1/2", "Баллы 300/220/100, выше 2 — 0"),
            (
                "Пустые поля, Unicode-цифры в паспорте, NaN, нулевой доход, невозможный стаж",
                "Ошибки формы; обе части анкеты не сохраняются частично",
            ),
            (
                "Порог предложения равен рейтингу/сроку/сумме/нагрузке",
                "Граница включается; неактивные предложения исключаются",
            ),
            ("Чужой идентификатор оценки", "404; данные другого пользователя не раскрываются"),
            (
                "Изменение анкеты или каталога после расчёта",
                "Исторический снимок прежний; новый расчёт создаёт новую запись",
            ),
            (
                "POST без CSRF и внешний next при входе",
                "CSRF-запрос отклонён; внешнее перенаправление не выполняется",
            ),
        ],
        [8, 9],
    )
    r.h("3 Соответствие задачам исходного плана")
    r.p(
        "Маркеры requirement связывают тесты с идентификаторами ЛР3. Список собирается при collection в test-map.json; tools/collect_evidence.py проверяет наличие тестов либо отдельного доказательства для каждой из 30 задач. Один сценарий может подтверждать несколько связанных задач; количество привязок не является количеством уникальных тестов."
    )
    rows = []
    for task in tasks:
        files = sorted({test.split("::")[0].removeprefix("tests/") for test in task["tests"]})
        evidence = ", ".join(files) if files else "; ".join(task["artifacts"])
        rows.append((task["id"], task["title"], evidence))
    r.table(
        "Матрица подтверждения всех запланированных задач",
        ["ID", "Планируемый результат", "Тесты / доказательства"],
        rows,
        [2.2, 7.5, 7.3],
    )
    r.h("4 Автоматические отчёты и результаты")
    stages = ["assessment-browser", "postgres-regression", "release"]
    rows = []
    for stage in stages:
        path = EVIDENCE / f"metrics/{stage}.json"
        if path.exists():
            m = json.loads(path.read_text())
            rows.append(
                (
                    stage,
                    m["database"],
                    m["tests"]["passed"],
                    m["tests"]["failures"] + m["tests"]["errors"],
                    f"{m['line_coverage_percent']} / {m['branch_coverage_percent']}",
                    m["test_duration_seconds"],
                )
            )
    r.table(
        "Результаты полного прогона",
        ["Этап", "БД", "Успех", "Ошибки", "Строки / ветви, %", "Время, с"],
        rows,
        [4.1, 2.7, 1.6, 1.9, 3.8, 2.9],
    )
    r.figure("docs/images/infrastructure/tests.png", "Автоматический HTML-отчёт pytest")
    r.figure(
        "docs/images/infrastructure/coverage.png",
        "HTML-отчёт покрытия с детализацией по Python-модулям",
    )
    r.p(
        "Я настроил формирование JUnit XML, HTML со статусом каждого теста, отчётов coverage JSON/XML/HTML, карты требований и метрик. При ошибке браузерного сценария сохраняются trace и screenshot. GitHub Actions публикует результаты в артефакте quality-evidence."
    )
    r.p(
        "Я выполнил 91 автоматическую проверку, включая два сценария Chromium. Получил 100% покрытия строк и ветвей Python-модулей core в установленной области измерения."
    )
    r.h("5 Выявленные проблемы инфраструктуры")
    r.table(
        "Исправленные дефекты",
        ["ID", "Причина", "Исправление"],
        [(d["id"], d["cause"], d["fix"]) for d in read("defects.json")["defects"]],
        [1.8, 7.2, 8],
    )
    r.p(
        "Я зарегистрировал и исправил четыре дефекта инфраструктуры. Проверку в PR №32 создал для проверки запрета слияния при ошибке CI. HTML-отчёт открываю через HTTP: python -m http.server 8090 --bind 127.0.0.1 --directory artifacts/runs/release. Адрес отчёта: http://127.0.0.1:8090/tests.html."
    )
    r.finish()


def lab8(summary, tasks):
    r = Report(8, "Отчёт о релизе")
    r.h("1 Состав завершённой системы")
    r.p(
        "Я завершил CreditSoft 1.0.0: реализовал регистрацию и вход, анкету, финансовые сведения, проверку данных, расчёт долговой нагрузки и рейтинга, каталог, подбор предложений и просмотр результата. Приложение запускаю через Docker Compose по адресу http://127.0.0.1:8080/."
    )
    r.figure("docs/images/application/login-desktop.png", "Вход в CreditSoft")
    r.figure(
        "docs/images/application/questionnaire-desktop.png",
        "Ввод, редактирование и сохранение анкеты",
    )
    r.p(
        "Я реализовал хранение ФИО, даты рождения, реквизитов паспорта, работодателя и стажа. Финансовая часть содержит доход, текущие платежи, сумму и срок кредита. Сервер проверяет поля и атомарно сохраняет обе части. Данные доступны после повторного входа и перезагрузки."
    )
    r.figure(
        "docs/images/application/validation-desktop.png",
        "Сообщения о некорректных данных рядом с полями",
    )
    r.h("2 Расчёт, подбор и сохранение результатов")
    r.p(
        "Я рассчитываю нагрузку как отношение текущих платежей к доходу, умноженное на 100%. Рейтинг 0–1000 состоит из оценки нагрузки до 450 баллов, стажа до 250 и отношения суммы кредита к годовому доходу до 300. Для расчётов использую Decimal и сравниваю границы до округления."
    )
    r.table(
        "Правила рейтинга edu-1.0",
        ["Фактор", "Границы → баллы"],
        [
            ("Нагрузка", "≤10% → 450; ≤30% → 350; ≤50% → 200; ≤70% → 80; >70% → 0"),
            ("Стаж в месяцах", "≥36 → 250; ≥12 → 180; ≥6 → 100; ≥1 → 40; 0 → 0"),
            ("Сумма / (доход × 12)", "≤0,5 → 300; ≤1 → 220; ≤2 → 100; >2 → 0"),
        ],
        [5.5, 11.5],
    )
    r.p(
        "Я отбираю активные предложения по минимальному рейтингу, максимальной нагрузке, сумме и сроку. Границы включительны. Сортирую результат по ставке и идентификатору. Для дохода 100 000 руб., платежей 15 000 руб., суммы 300 000 руб. и стажа 48 месяцев получил нагрузку 15%, рейтинг 900 и четыре предложения."
    )
    r.figure(
        "docs/images/application/results-desktop.png",
        "Итоговая оценка, объяснение баллов и предложения",
    )
    r.figure(
        "docs/images/application/results-no-offers.png",
        "Корректное состояние при отсутствии подходящих предложений",
    )
    r.figure("docs/images/application/results-mobile.png", "Результат расчёта на мобильном экране")
    r.p(
        "Я сохраняю в Assessment входные финансовые параметры, баллы, версию правил и подобранные предложения. При изменении анкеты или каталога предыдущий результат сохраняется. Новый POST создаёт следующую оценку, GET читает готовую. Доступ к чужому результату возвращает 404."
    )
    r.h("3 Компоненты и последовательность вызовов")
    r.table(
        "Реализованные части системы",
        ["Компонент", "Назначение"],
        [
            (
                "HTML/CSS/JavaScript + Bootstrap",
                "Формы, сообщения, карточки предложений и адаптивное представление",
            ),
            ("accounts.py / Django Auth", "Регистрация, хеширование пароля, вход, сессии и выход"),
            (
                "questionnaire.py / forms.py / validation.py",
                "Обработка анкеты, проверка и координация сохранения",
            ),
            (
                "repository.py / models.py",
                "ORM, транзакции, связь пользователя с анкетой, каталогом и оценками",
            ),
            ("scoring.py", "Чистый расчёт нагрузки, баллов и текстовых объяснений"),
            ("offers.py", "Отбор активных предложений по всем ограничениям"),
            ("results.py", "Создание снимка и выдача результатов владельцу"),
            (
                "Nginx → Gunicorn/Django → PostgreSQL",
                "Раздача статики и проксирование → выполнение Python → долговременное хранение",
            ),
        ],
        [6.5, 10.5],
    )
    r.p(
        "Я организовал обработку запроса следующим образом: браузер → Nginx → Gunicorn → Django. Session middleware определяет пользователя; questionnaire читает анкету через repository и ORM. При сохранении сервер проверяет обе формы и записывает данные транзакцией. При расчёте results блокирует строку пользователя, читает анкету, вызывает scoring и offers, сохраняет Assessment и перенаправляет на результат. Повторное открытие страницы читает сохранённый снимок."
    )
    r.h("4 Готовность задач и дефекты")
    r.table(
        "Сводка GitHub Issues",
        ["Категория", "Всего", "Завершено"],
        [
            ("Исходные функциональные задачи CR", 24, summary["closed_counts"].get("original", 0)),
            (
                "Дополнительные задачи плана ЛР3",
                6,
                summary["closed_counts"].get("planned_additional", 0),
            ),
            ("Задачи-описания фактических дефектов", summary["defects"], summary["closed_defects"]),
            (
                "Всего",
                30 + summary["defects"],
                sum(summary["closed_counts"].values()) + summary["closed_defects"],
            ),
        ],
        [11, 3, 3],
    )
    r.p(
        "Я выполнил 24 функциональные задачи CR и шесть дополнительных задач плана ЛР3: PR-01, INF-01, QA-01, RL-01, QA-02, RL-02. Зарегистрировал и исправил BUG-01–BUG-04. Состояния задач получил через GitHub API и сохранил в docs/evidence/github-export.json; соответствие задач проверкам — в acceptance.json."
    )
    r.figure("docs/images/infrastructure/tasks.png", "Закрытые задачи в GitHub Issues")
    r.table(
        "Выполнение задач и исходные оценки",
        ["ID", "Задача", "План, ч", "Состояние"],
        [
            (
                t["id"],
                t["title"],
                t["planned_hours"],
                "Закрыта" if t["state"] == "closed" else "Открыта",
            )
            for t in tasks
        ],
        [2.2, 10.5, 2, 2.3],
    )
    r.h("5 Фактическое время и отклонения")
    r.p(
        "В плане ЛР3 я оценил 30 задач в 140 часов: 122 часа функциональных работ и 18 часов дополнительных. Для контроля этапов использовал tools/tasklog.py, фиксирующий начало и окончание в UTC. В таблице привёл интервалы между этими отметками, включая ожидание CI и паузы."
    )
    plan = {
        "infrastructure": 6,
        "prototype": sum(
            t["planned_hours"]
            for t in tasks
            if t["id"] == "PR-01" or t["id"].startswith(tuple(f"CR-0{n}" for n in range(1, 5)))
        ),
        "assessment": sum(
            t["planned_hours"]
            for t in tasks
            if t["id"].startswith(tuple(f"CR-0{n}" for n in range(5, 9)))
        ),
        "validation": 6,
        "release": 2,
    }
    labels = {
        "infrastructure": "Окружение, INF-01 / RL-01",
        "prototype": "Прототип, PR-01 и CR-01–04",
        "assessment": "Основной функционал, CR-05–08",
        "validation": "Проверки, QA-01 / QA-02",
        "release": "Документация и выпуск, RL-02",
    }
    rows = [
        (
            labels[p["phase"]],
            plan[p["phase"]],
            f"{p['elapsed_seconds'] / 60:.2f}",
            "Идёт на момент снимка" if p.get("snapshot_in_progress") else "Завершён",
        )
        for p in summary["worklog"]
    ]
    r.table(
        "Измеренное время по этапам",
        ["Этап и задачи", "План, ч", "Интервал, мин", "Учёт"],
        rows,
        [8.2, 2, 3, 3.8],
    )
    r.p(
        f"Я зафиксировал сумму интервалов этапов: {sum(p['elapsed_seconds'] for p in summary['worklog']) / 60:.2f} мин. Плановые оценки сгруппировал без повторного учёта: INF-01 и RL-01 отнёс к инфраструктуре, QA-01 и QA-02 — к проверкам, RL-02 — к выпуску."
    )
    r.table(
        "Причины отличий интервалов от плановых оценок",
        ["Этап", "Причина"],
        [
            (
                "Инфраструктура",
                "Я использовал готовые Docker-образы. В интервал вошли загрузка образов и ожидание CI.",
            ),
            (
                "Прототип",
                "Я использовал Django Auth, ORM, формы и Bootstrap.",
            ),
            (
                "Основной функционал",
                "Я реализовал правила скоринга и каталог внутри приложения.",
            ),
            (
                "Проверки",
                "Я использовал параметризацию и автоматические прогоны; исправил ошибки инфраструктуры.",
            ),
            (
                "Выпуск",
                "Я настроил экспорт PDF через LibreOffice после зависания Word. Интервал содержит паузу между этапами работы.",
            ),
        ],
        [4.3, 12.7],
    )
    r.h("6 Закрытие задач, активность и объём проекта")
    r.figure(
        "docs/images/statistics/burndown.png",
        "Плановое и фактическое сгорание оценок задач по неделям",
    )
    r.p(
        "Я построил плановую линию по восьми двухнедельным спринтам ЛР3 от 07.09.2026 до 27.12.2026. Закрытие в плане отнёс на конец спринта. Для фактической линии вычитал оценку задачи в дату closed_at из GitHub. График показывает остаток плановых часов."
    )
    r.figure(
        "docs/images/statistics/weekly-activity.png",
        "Закрытые карточки, коммиты и изменения строк по неделям",
    )
    measured = [w for w in summary["weeks"] if w["commits"] is not None]
    r.table(
        "Наблюдаемая недельная статистика",
        ["Начало недели", "Закрыто карточек", "Коммитов без merge", "+ строк", "− строк"],
        [
            (w["week"], w["closed_issues"], w["commits"], w["added_lines"], w["removed_lines"])
            for w in measured
        ],
        [3.5, 3.7, 4, 2.9, 2.9],
    )
    r.p(
        "Я сгруппировал закрытые задачи по датам GitHub, а коммиты — по датам Git. В статистику включил коммиты, достижимые из HEAD, без merge-коммитов и отдельной контрольной ветки ci/protection-probe."
    )
    v = summary["volume"]
    r.table(
        "Объём репозитория на момент снимка",
        ["Показатель", "Значение"],
        [
            ("Файлы под управлением Git", v["tracked_files"]),
            ("Собственные файлы исходников, тестов и конфигураций", v["authored_source_files"]),
            ("Физические строки этих файлов", v["authored_physical_lines"]),
            ("HEAD статистики", summary["head"]),
            ("Момент выгрузки UTC", summary["recorded_at"]),
        ],
        [10.2, 6.8],
    )
    r.p(
        "Я подсчитал файлы .py, .html, .css, .js, .yaml и .yml, включая tests и tools. Исключил Bootstrap и сгенерированные миграции. Физические строки включают пустые строки и комментарии. Состав репозитория по расширениям сохранил в summary.json."
    )
    r.h("7 Проверка выпуска и эксплуатация")
    r.p(
        "Я проверил систему 91 тестом на SQLite и PostgreSQL, включая два сценария Chromium. Проверил запуск Docker-стенда через Nginx/Gunicorn и повторный вход после перезапуска. В smoke-проверке получил рейтинг 900 и четыре предложения. Создал резервную копию pg_dump и проверил её оглавление командой pg_restore --list. Отчёты, пакеты и контрольные суммы разместил в GitHub Releases v1.0.0."
    )
    r.p(
        "Я подготовил запуск: python tools/bootstrap.py, затем docker compose up -d --build. Проверку доступности выполняю через /health/ и python tools/smoke.py, полный набор проверок — python tools/check.py --stage local --e2e. Данные PostgreSQL сохраняются в именованном томе, секреты — в .local."
    )
    r.finish()


def main():
    summary = read("summary.json")
    tasks = read("acceptance.json")
    for builder in [lab5, lab6, lab7, lab8]:
        builder(summary, tasks)


if __name__ == "__main__":
    main()
